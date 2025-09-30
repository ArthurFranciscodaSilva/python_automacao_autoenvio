import pandas as pd
from openpyxl import load_workbook
from pathlib import Path

# ==== CONFIGURAÇÕES ====

# Diretório onde estão os arquivos baixados do AutoEnvio
DIRETORIO_AUTOENVIO = Path(r"C:\Users\Arthur\Downloads\AutoEnvio\TETO\30.09.25.2")

# Caminho do arquivo Excel (base) no OneDrive (via sincronização local)
EXCEL_ONEDRIVE = Path(r"C:\Users\Arthur\OneDrive\Área de Trabalho\DISPAROS WHATSAPP\NÃO RETRABALHAR DISPAROS\DIRETO_AUTOENVIO_TETO.xlsx")

# Nome da aba da planilha base
NOME_ABA = "Plan1"

# Nome das colunas a serem copiadas e inseridas
COLUNA_ORIGEM_NOME = "Nome"
COLUNA_ORIGEM_NUMERO = "Número"
COLUNA_DESTINO_NOME = "NOME"
COLUNA_DESTINO_NUMERO = "NÚMERO"

# ==== COLETAR DADOS DOS ARQUIVOS AUTOENVIO ====

linhas_coletadas = []

for arquivo in DIRETORIO_AUTOENVIO.glob("*.xlsx"):
    try:
        df = pd.read_excel(arquivo, dtype=str)
    except Exception as e:
        print(f"Erro ao ler {arquivo.name}: {e}")
        continue

    filtro = df["Erro"] == "O contato não é um número de WhatsApp"
    for _, linha in df.loc[filtro].iterrows():
        nome = str(linha.get(COLUNA_ORIGEM_NOME, "")).strip()
        numero = str(linha.get(COLUNA_ORIGEM_NUMERO, "")).strip()
        if nome and numero:
            linhas_coletadas.append((nome, numero))

print(f"Total de registros coletados: {len(linhas_coletadas)}")

# ==== ABRIR PLANILHA BASE ====

wb = load_workbook(EXCEL_ONEDRIVE)
if NOME_ABA not in wb.sheetnames:
    raise ValueError(f"Aba '{NOME_ABA}' não encontrada na planilha base.")

ws = wb[NOME_ABA]

# Mapear as colunas pelo cabeçalho
colunas = {cell.value: cell.column for cell in ws[1]}

if COLUNA_DESTINO_NOME not in colunas:
    raise ValueError(f"Coluna destino '{COLUNA_DESTINO_NOME}' não encontrada.")
if COLUNA_DESTINO_NUMERO not in colunas:
    raise ValueError(f"Coluna destino '{COLUNA_DESTINO_NUMERO}' não encontrada.")

col_dest_nome = colunas[COLUNA_DESTINO_NOME]
col_dest_numero = colunas[COLUNA_DESTINO_NUMERO]

# ==== INSERIR DADOS ====

linha_atual = ws.max_row + 1

for nome, numero in linhas_coletadas:
    ws.cell(row=linha_atual, column=col_dest_nome).value = nome
    ws.cell(row=linha_atual, column=col_dest_numero).value = numero
    linha_atual += 1

# ==== SALVAR ALTERAÇÕES ====

wb.save(EXCEL_ONEDRIVE)
print("Planilha base atualizada com sucesso!")
